"""AI matnidan .pptx yig'ish: tayyor mavzular, admin yuklagan shablonlar va rasm manbalari."""
import asyncio
import glob
import io
import os
import random
import aiohttp
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
import db
from config import TEMPLATE_DIR, UNSPLASH_KEY

UA = {"User-Agent": "ZehnAIBot/1.0 (educational student bot)"}
WHITE = (0xFF, 0xFF, 0xFF)
THEMES = {
    "🔵 Tibbiyot": dict(main=(0x0B, 0x3C, 0x7D), accent=(0x2E, 0x9B, 0xD6), light=(0xE8, 0xF1, 0xFB), text=(0x22, 0x2B, 0x38)),
    "🟢 Yashil": dict(main=(0x0F, 0x6B, 0x4F), accent=(0x3C, 0xB3, 0x71), light=(0xE6, 0xF5, 0xEE), text=(0x1F, 0x2D, 0x27)),
    "🟣 Binafsha": dict(main=(0x4B, 0x2A, 0x8C), accent=(0x9B, 0x6B, 0xE0), light=(0xF0, 0xEA, 0xFB), text=(0x2A, 0x20, 0x3D)),
    "🔴 Qizil": dict(main=(0xA3, 0x1D, 0x2B), accent=(0xE0, 0x5A, 0x5A), light=(0xFB, 0xEA, 0xEA), text=(0x3A, 0x1F, 0x22)),
    "⚫ Qora-oltin": dict(main=(0x1F, 0x24, 0x2E), accent=(0xF2, 0xB1, 0x34), light=(0xF3, 0xF4, 0xF6), text=(0x1F, 0x24, 0x2E)),
}


# ---------------- shablonlar ----------------
def _file_templates():
    return {os.path.splitext(os.path.basename(p))[0]: p for p in glob.glob(os.path.join(TEMPLATE_DIR, "*.pptx"))
            if not os.path.basename(p).startswith("~$")}


async def template_names():
    names = list(_file_templates())
    names += [a["name"] for a in await db.assets("template") if a["name"] not in names]
    return names


async def load_template(bot, name):
    """Shablon baytlarini qaytaradi (papkadan yoki admin yuklagan fayldan)."""
    path = _file_templates().get(name)
    if path:
        with open(path, "rb") as f:
            return f.read()
    for a in await db.assets("template"):
        if a["name"] == name:
            buf = await bot.download(a["file_id"])
            return buf.read()
    return None


# ---------------- rasmlar ----------------
async def _get(session, url, **kw):
    async with session.get(url, headers=UA, **kw) as r:
        r.raise_for_status()
        return await r.content.read(5_000_000) if kw.get("params") is None else await r.json()


async def _library(bot, lib, hay):
    for a in lib:
        kws = [k.strip().lower() for k in (a["keywords"] or "").split(",") if k.strip()]
        if any(k in hay for k in kws):
            try:
                buf = await bot.download(a["file_id"])
                return buf.read(), "Kutubxona"
            except Exception:
                continue
    return None


async def _wikimedia(session, query):
    if not query:
        return None
    try:
        data = await _get(session, "https://commons.wikimedia.org/w/api.php", params={
            "action": "query", "generator": "search", "gsrsearch": f"{query} filetype:bitmap", "gsrnamespace": 6,
            "gsrlimit": 5, "prop": "imageinfo", "iiprop": "url|mime", "iiurlwidth": 900, "format": "json"})
        pages = sorted(data["query"]["pages"].values(), key=lambda p: p.get("index", 99))
        for p in pages:
            ii = p["imageinfo"][0]
            if ii.get("mime") in ("image/jpeg", "image/png"):
                return await _get(session, ii["thumburl"]), "Wikimedia Commons"
    except Exception:
        pass
    return None


async def _unsplash(session, query):
    if not UNSPLASH_KEY or not query:
        return None
    try:
        async with session.get("https://api.unsplash.com/search/photos", headers={"Authorization": f"Client-ID {UNSPLASH_KEY}"},
                               params={"query": query, "per_page": 1, "orientation": "landscape"}) as r:
            data = await r.json()
        return await _get(session, data["results"][0]["urls"]["small"]), "Unsplash"
    except Exception:
        return None


async def collect_images(data, bot):
    """{slayd_indeksi: (rasm_baytlari, manba)}. Tartib: kutubxona -> Wikimedia -> Unsplash."""
    out = {}
    lib = await db.assets("image")
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=20)) as session:
        async def one(i, sd):
            q = sd.get("image_query", "")
            hay = f"{q} {sd.get('title', '')} {' '.join(map(str, sd.get('bullets', [])))}".lower()
            res = await _library(bot, lib, hay) or await _wikimedia(session, q) or await _unsplash(session, q)
            if res:
                out[i] = res
        await asyncio.gather(*[one(i, sd) for i, sd in enumerate(data["slides"][:10])])
    return out


# ---------------- yig'ish ----------------
def _rgb(t):
    return RGBColor(*t)


def _rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(l), int(t), int(w), int(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    return s


def _text(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(int(l), int(t), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, _rgb(color)
    return tb


def _bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = _rgb(color)


def _fit_picture(slide, img, bx, by, bw, bh):
    try:
        pic = slide.shapes.add_picture(io.BytesIO(img), int(bx), int(by), width=int(bw))
        if pic.height > bh:
            r = bh / pic.height
            pic.width, pic.height = int(pic.width * r), int(bh)
        pic.left = int(bx + (bw - pic.width) / 2)
        pic.top = int(by + (bh - pic.height) / 2)
    except Exception:
        pass


def _builtin(prs, data, author, images, th):
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    main, acc, light, txt = th["main"], th["accent"], th["light"], th["text"]
    # titul
    s = prs.slides.add_slide(blank)
    _bg(s, main)
    _rect(s, 0, 0, Inches(0.4), H, acc)
    _text(s, Inches(1.2), Inches(1.6), W - Inches(2.4), Inches(2.6), data.get("title", ""), 44, WHITE, True,
          anchor=MSO_ANCHOR.BOTTOM)
    _rect(s, Inches(1.25), Inches(4.4), Inches(2.6), Inches(0.09), acc)
    _text(s, Inches(1.2), Inches(4.7), W - Inches(2.4), Inches(1.4), author, 20, light)
    _text(s, Inches(1.2), H - Inches(0.9), Inches(4), Inches(0.4), "Zehn AI", 14, acc, True)
    total = len(data["slides"])
    for i, sd in enumerate(data["slides"]):
        s = prs.slides.add_slide(blank)
        _bg(s, WHITE)
        _rect(s, 0, 0, W, Inches(1.25), main)
        _rect(s, 0, Inches(1.25), W, Inches(0.07), acc)
        _text(s, Inches(0.6), 0, W - Inches(1.2), Inches(1.25), sd.get("title", ""), 32, WHITE, True,
              anchor=MSO_ANCHOR.MIDDLE)
        cx, cy, cw, ch = Inches(0.5), Inches(1.7), W - Inches(1.0), H - Inches(2.5)
        _rect(s, cx, cy, cw, ch, light)
        img = images.get(i)
        tw = cw - Inches(0.8) - (Inches(5.3) if img else 0)
        bullets = [str(b) for b in sd.get("bullets", [])][:6]
        size = 24 if sum(map(len, bullets)) < 300 else 20 if sum(map(len, bullets)) < 480 else 17
        tb = s.shapes.add_textbox(int(cx + Inches(0.4)), int(cy + Inches(0.3)), int(tw), int(ch - Inches(0.6)))
        tf = tb.text_frame
        tf.word_wrap = True
        for k, b in enumerate(bullets):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            r1 = p.add_run()
            r1.text = "■  "
            r1.font.size, r1.font.color.rgb = Pt(size * 0.6), _rgb(acc)
            r2 = p.add_run()
            r2.text = b
            r2.font.size, r2.font.color.rgb = Pt(size), _rgb(txt)
        if img:
            _fit_picture(s, img[0], cx + cw - Inches(5.1), cy + Inches(0.3), Inches(4.7), ch - Inches(0.6))
            s.notes_slide.notes_text_frame.text = f"Rasm manbasi: {img[1]}"
        _text(s, Inches(0.6), H - Inches(0.65), Inches(4), Inches(0.4), "Zehn AI", 12, acc, True)
        _text(s, W - Inches(3.2), H - Inches(0.65), Inches(2.7), Inches(0.4), f"{i + 1} / {total}", 12,
              (0x88, 0x92, 0xA0), align=PP_ALIGN.RIGHT)


def _body(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            return ph
    return None


def _custom(prs, data, author, images):
    lst = prs.slides._sldIdLst          # shablondagi namunaviy slaydlarni o'chiramiz
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)
    W, H = prs.slide_width, prs.slide_height
    lay_t = prs.slide_layouts[0]
    lay_b = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else lay_t
    s = prs.slides.add_slide(lay_t)
    if s.shapes.title is not None:
        s.shapes.title.text = data.get("title", "")
    sub = _body(s)
    if sub is not None:
        sub.text = author
    for i, sd in enumerate(data["slides"]):
        s = prs.slides.add_slide(lay_b)
        if s.shapes.title is not None:
            s.shapes.title.text = sd.get("title", "")
        body = _body(s)
        bullets = [str(b) for b in sd.get("bullets", [])][:6]
        if body is not None and bullets:
            tf = body.text_frame
            tf.word_wrap = True
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b
        img = images.get(i)
        if img:
            if body is not None:
                body.width = int(W * 0.5)
            _fit_picture(s, img[0], W * 0.57, H * 0.25, W * 0.38, H * 0.6)
            s.notes_slide.notes_text_frame.text = f"Rasm manbasi: {img[1]}"


def build(data, author, out_path, images, theme=None, template_bytes=None):
    if template_bytes:
        prs = Presentation(io.BytesIO(template_bytes))
        _custom(prs, data, author, images)
    else:
        prs = Presentation()
        _builtin(prs, data, author, images, THEMES.get(theme) or random.choice(list(THEMES.values())))
    prs.save(out_path)
